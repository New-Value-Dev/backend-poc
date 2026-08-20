"""승인된 proofread finding을 원본 파일에 반영해 새 버전을 만든다.

포맷별 담당:
- DOCX/HWPX/PPTX: 실제 텍스트 편집. run/텍스트노드 경계를 인식하는 부분 치환으로
  구간 밖 서식은 건드리지 않는다 (_replace_span_in_nodes). 표 섹션은 원문이 들어있는
  셀 하나만 찾아 그 안에서 치환한다(_apply_*_table_cell) — 셀 경계를 넘어서는
  span은 애초에 찾지 않으므로 다른 셀 데이터가 섞일 일이 없다.
- PDF: 완전 편집이 불가능하므로 bbox 자리를 redact 후 텍스트를 다시 그려 넣는
  시각적 패치. 다시 그려 넣은 텍스트도 실제 텍스트 레이어라 재파싱하면 반영된다.
- TXT/MD: 서식 구조가 없어 블록 단위 문자열 치환으로 충분하다.
"""

from __future__ import annotations

import io
import logging
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

import fitz
from docx import Document as load_docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import Depends
from pptx import Presentation
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from app.core.config import Settings, get_settings
from app.core.database import get_db
from app.models.document_section import DocumentSection
from app.models.document_version import DocumentVersion
from app.repositories.analysis_repository import AiAnalysisResultRepository
from app.repositories.document_repository import DocumentVersionRepository
from app.schemas.proofread import FindingStatus, ProofreadFindingRead
from app.services.document_service import (
    DocumentService,
    VersionNotFoundError,
    get_document_service,
)
from app.services.parsing.base import ParseError
from app.services.parsing.docx_parser import iter_block_items
from app.services.parsing.docx_parser import table_text as docx_table_text
from app.services.parsing.hwpx_parser import SECTION_RE, iter_body_blocks, iter_t_nodes, iter_table_cells
from app.services.parsing.hwpx_parser import table_text as hwpx_table_text
from app.services.parsing.hwpx_parser import text_excluding_tables as hwpx_text
from app.services.parsing.pptx_parser import iter_slide_blocks
from app.services.parsing.pptx_parser import table_text as pptx_table_text
from app.services.parsing.text_parser import BLANK_LINE_RE
from app.services.proofread_service import AnalysisNotCompletedError, AnalysisNotFoundError, rehydrate_findings
from app.services.storage import LocalStorageProvider, StorageProvider

logger = logging.getLogger(__name__)

_MAX_SKIPPED_REASONS = 20


class NoAcceptedFindingsError(Exception):
    def __init__(self, analysis_id: int) -> None:
        self.analysis_id = analysis_id
        super().__init__(f"Analysis {analysis_id} has no applicable accepted findings")


class UnsupportedFileTypeError(Exception):
    def __init__(self, ext: str) -> None:
        self.ext = ext
        super().__init__(f"'.{ext}' 파일은 교정 적용을 지원하지 않습니다")


class AlreadyAppliedError(Exception):
    def __init__(self, analysis_id: int) -> None:
        self.analysis_id = analysis_id
        super().__init__(f"Analysis {analysis_id} has already been applied")


@dataclass
class ApplyOutcome:
    new_version: DocumentVersion
    applied_count: int
    skipped_count: int
    skipped_reasons: list[str]


def _replace_span_in_nodes(nodes: list, original: str, suggestion: str) -> bool:
    """nodes를 순서대로 이어붙인 전체 텍스트에서 original을 찾아 suggestion으로 바꾼다.

    구간과 겹치지 않는 노드는 절대 건드리지 않으므로, DOCX run이나 HWPX <hp:t>처럼
    문단이 여러 서식 단위로 쪼개져 있어도 구간 밖 서식이 보존된다. node는 .text
    getter/setter만 있으면 되므로(python-docx Run, ElementTree Element 둘 다 해당)
    포맷에 관계없이 재사용한다.
    """
    texts = [n.text or "" for n in nodes]
    full_text = "".join(texts)
    idx = full_text.find(original)
    if idx == -1:
        return False
    end = idx + len(original)

    pos = 0
    replaced_first = False
    for node, text in zip(nodes, texts):
        nstart, nend = pos, pos + len(text)
        pos = nend
        if nend <= idx or nstart >= end:
            continue  # 이 span과 겹치지 않는 노드 — 그대로 둠
        local_start = max(0, idx - nstart)
        local_end = min(len(text), end - nstart)
        prefix, suffix = text[:local_start], text[local_end:]
        if not replaced_first:
            node.text = prefix + suggestion + suffix
            replaced_first = True
        else:
            node.text = prefix + suffix  # 겹친 나머지 부분만 제거
    return True


def _apply_docx_table_cell(table: Table, original: str, suggestion: str) -> bool:
    """표 전체가 아니라 원문을 담고 있는 셀 하나만 찾아 그 안에서 치환한다.

    셀 경계를 넘어 span을 찾으면 엉뚱한 셀 데이터가 섞일 수 있어, 셀 단위로 먼저
    좁혀서 _replace_span_in_nodes 를 적용한다(첫 매치만 반영, 나머지는 그대로 둠).
    """
    for row in table.rows:
        for cell in row.cells:
            nodes = [run for paragraph in cell.paragraphs for run in paragraph.runs]
            if _replace_span_in_nodes(nodes, original, suggestion):
                return True
    return False


def _apply_hwpx_table_cell(tbl_elem: ET.Element, original: str, suggestion: str) -> bool:
    for tc in iter_table_cells(tbl_elem):
        if _replace_span_in_nodes(list(iter_t_nodes(tc)), original, suggestion):
            return True
    return False


def _apply_pptx_table_cell(table, original: str, suggestion: str) -> bool:
    for row in table.rows:
        for cell in row.cells:
            nodes = [run for paragraph in cell.text_frame.paragraphs for run in paragraph.runs]
            if _replace_span_in_nodes(nodes, original, suggestion):
                return True
    return False


def _group_by_section(
    accepted: list[ProofreadFindingRead], skipped: list[str]
) -> dict[int, list[ProofreadFindingRead]]:
    by_section: dict[int, list[ProofreadFindingRead]] = defaultdict(list)
    for f in accepted:
        if f.section_id is None:
            skipped.append(f"finding {f.id}: section_id 없음")
            continue
        by_section[f.section_id].append(f)
    return by_section


def _expected_text(section: DocumentSection) -> str:
    return (section.title if section.section_type == "heading" else section.content) or ""


def _apply_docx(
    path: Path, accepted: list[ProofreadFindingRead], sections: dict[int, DocumentSection]
) -> tuple[bytes, int, list[str]]:
    document = load_docx(str(path))
    blocks: list[Paragraph | Table] = []
    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            if block.text.strip():
                blocks.append(block)
        elif isinstance(block, Table):
            if docx_table_text(block).strip():
                blocks.append(block)

    applied = 0
    skipped: list[str] = []
    by_section = _group_by_section(accepted, skipped)

    for section_id, section_findings in by_section.items():
        section = sections.get(section_id)
        if section is None:
            skipped.extend(f"finding {f.id}: section {section_id} 없음" for f in section_findings)
            continue
        if section.order_no >= len(blocks):
            skipped.extend(
                f"finding {f.id}: order {section.order_no} 범위 밖" for f in section_findings
            )
            continue
        block = blocks[section.order_no]

        if isinstance(block, Table):
            if docx_table_text(block).strip() != _expected_text(section).strip():
                skipped.extend(
                    f"finding {f.id}: 표 텍스트가 분석 시점과 달라짐" for f in section_findings
                )
                continue
            for f in section_findings:
                if _apply_docx_table_cell(block, f.original, f.suggestion):
                    applied += 1
                else:
                    skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")
            continue

        if block.text.strip() != _expected_text(section).strip():
            skipped.extend(
                f"finding {f.id}: 문단 텍스트가 분석 시점과 달라짐" for f in section_findings
            )
            continue

        for f in section_findings:
            if _replace_span_in_nodes(block.runs, f.original, f.suggestion):
                applied += 1
            else:
                skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue(), applied, skipped


def _apply_hwpx(
    path: Path, accepted: list[ProofreadFindingRead], sections: dict[int, DocumentSection]
) -> tuple[bytes, int, list[str]]:
    applied = 0
    skipped: list[str] = []
    by_section = _group_by_section(accepted, skipped)

    with zipfile.ZipFile(path) as zf:
        names = sorted(n for n in zf.namelist() if SECTION_RE.search(n))
        if not names:
            raise ParseError("HWPX 본문(Contents/section*.xml)을 찾을 수 없습니다")
        trees: dict[str, ET.Element] = {name: ET.fromstring(zf.read(name)) for name in names}
        # HWPX(ODF/EPUB 계열과 동일하게)는 zip의 첫 엔트리 mimetype 이 무압축(STORED)이어야 하는
        # 컨테이너 규약이 있다. infolist() 로 원본 순서 + 엔트리별 compress_type 을 그대로 들고 있다가
        # 쓸 때 재사용해야지, 전체를 ZIP_DEFLATED 로 새로 쓰면 이 규약이 깨져 파일이 안 열린다.
        infos = zf.infolist()
        original_bytes = {info.filename: zf.read(info.filename) for info in infos}

    # section.order 는 section*.xml 파일 전체를 통틀어 이어지는 전역 인덱스
    # (HwpxParser.parse() 참고) — 그대로 재현해서 몇 번째 파일의 어느 엘리먼트인지 기록.
    blocks: list[tuple[str, ET.Element, str]] = []
    for name in names:
        for kind, elem, _text in iter_body_blocks(trees[name]):
            blocks.append((name, elem, kind))

    modified_files: set[str] = set()
    for section_id, section_findings in by_section.items():
        section = sections.get(section_id)
        if section is None:
            skipped.extend(f"finding {f.id}: section {section_id} 없음" for f in section_findings)
            continue
        if section.order_no >= len(blocks):
            skipped.extend(
                f"finding {f.id}: order {section.order_no} 범위 밖" for f in section_findings
            )
            continue
        name, elem, kind = blocks[section.order_no]

        if kind == "tbl":
            if hwpx_table_text(elem).strip() != _expected_text(section).strip():
                skipped.extend(
                    f"finding {f.id}: 표 텍스트가 분석 시점과 달라짐" for f in section_findings
                )
                continue
            for f in section_findings:
                if _apply_hwpx_table_cell(elem, f.original, f.suggestion):
                    applied += 1
                    modified_files.add(name)
                else:
                    skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")
            continue

        if hwpx_text(elem).strip() != _expected_text(section).strip():
            skipped.extend(
                f"finding {f.id}: 문단 텍스트가 분석 시점과 달라짐" for f in section_findings
            )
            continue

        t_nodes = list(iter_t_nodes(elem))
        for f in section_findings:
            if _replace_span_in_nodes(t_nodes, f.original, f.suggestion):
                applied += 1
                modified_files.add(name)
            else:
                skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as out:
        for info in infos:
            name = info.filename
            if name in modified_files:
                data = ET.tostring(trees[name], encoding="utf-8", xml_declaration=True)
            else:
                data = original_bytes[name]
            # info(ZipInfo)를 그대로 재사용해 원본 compress_type/순서를 보존한다.
            out.writestr(info, data)
    return buf.getvalue(), applied, skipped


def _apply_pptx(
    path: Path, accepted: list[ProofreadFindingRead], sections: dict[int, DocumentSection]
) -> tuple[bytes, int, list[str]]:
    prs = Presentation(str(path))
    blocks: list[tuple[str, object]] = [
        (kind, obj) for kind, obj, _text, _slide_index in iter_slide_blocks(prs)
    ]

    applied = 0
    skipped: list[str] = []
    by_section = _group_by_section(accepted, skipped)

    for section_id, section_findings in by_section.items():
        section = sections.get(section_id)
        if section is None:
            skipped.extend(f"finding {f.id}: section {section_id} 없음" for f in section_findings)
            continue
        if section.order_no >= len(blocks):
            skipped.extend(
                f"finding {f.id}: order {section.order_no} 범위 밖" for f in section_findings
            )
            continue
        kind, obj = blocks[section.order_no]

        if kind == "table":
            table = obj
            if pptx_table_text(table).strip() != _expected_text(section).strip():
                skipped.extend(
                    f"finding {f.id}: 표 텍스트가 분석 시점과 달라짐" for f in section_findings
                )
                continue
            for f in section_findings:
                if _apply_pptx_table_cell(table, f.original, f.suggestion):
                    applied += 1
                else:
                    skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")
            continue

        paragraph = obj
        if paragraph.text.strip() != _expected_text(section).strip():
            skipped.extend(
                f"finding {f.id}: 문단 텍스트가 분석 시점과 달라짐" for f in section_findings
            )
            continue

        for f in section_findings:
            if _replace_span_in_nodes(paragraph.runs, f.original, f.suggestion):
                applied += 1
            else:
                skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), applied, skipped


_MIN_PDF_FONT_SIZE = 6.0


def _insert_textbox_fit(
    page, rect: "fitz.Rect", text: str, fontsize: float, fontname: str = "korea-s", color=(0, 0, 0)
) -> None:
    """rect 에 안 들어가면 폰트를 줄여가며 재시도하고, 최소 크기에서도 안 들어가면
    rect 세로를 늘려서라도 반드시 그린다.

    `insert_textbox` 는 텍스트가 rect 를 넘치면(음수 반환) **아무것도 그리지 않고
    조용히 실패**한다 — 반환값을 안 보면 알 수가 없다. 이미 `add_redact_annot` 로
    원문을 지운 뒤라, 이 실패를 그대로 두면 그 자리가 통째로 빈 칸이 된다.
    """
    size = max(fontsize, _MIN_PDF_FONT_SIZE)
    rc = page.insert_textbox(rect, text, fontsize=size, fontname=fontname, color=color)
    while rc < 0 and size > _MIN_PDF_FONT_SIZE:
        size = max(size - 0.5, _MIN_PDF_FONT_SIZE)
        rc = page.insert_textbox(rect, text, fontsize=size, fontname=fontname, color=color)
    if rc < 0:
        grown = fitz.Rect(rect.x0, rect.y0, rect.x1, rect.y1 - rc + 2)
        page.insert_textbox(grown, text, fontsize=size, fontname=fontname, color=color)


def _apply_pdf(
    path: Path, accepted: list[ProofreadFindingRead], sections: dict[int, DocumentSection]
) -> tuple[bytes, int, list[str]]:
    doc = fitz.open(str(path))
    applied = 0
    skipped: list[str] = []
    by_section = _group_by_section(accepted, skipped)

    try:
        for section_id, section_findings in by_section.items():
            section = sections.get(section_id)
            if section is None:
                skipped.extend(
                    f"finding {f.id}: section {section_id} 없음" for f in section_findings
                )
                continue

            meta = section.meta or {}
            bbox = meta.get("bbox")
            page_no = section.page_start
            if not bbox or page_no is None or not (1 <= page_no <= doc.page_count):
                skipped.extend(
                    f"finding {f.id}: bbox/page 정보 없음" for f in section_findings
                )
                continue

            page = doc[page_no - 1]
            rect = fitz.Rect(*bbox)
            expected = _expected_text(section)
            live_text = page.get_text("text", clip=rect).strip()
            if live_text and expected and live_text not in expected and expected not in live_text:
                skipped.extend(
                    f"finding {f.id}: 페이지 텍스트가 분석 시점과 달라짐" for f in section_findings
                )
                continue

            new_text = expected.strip()
            section_applied = 0
            for f in section_findings:
                if f.original in new_text:
                    new_text = new_text.replace(f.original, f.suggestion, 1)
                    section_applied += 1
                else:
                    skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")
            if section_applied == 0:
                continue

            font_size = meta.get("font_size") or meta.get("line_height") or 10.0
            page.add_redact_annot(rect, fill=(1, 1, 1))
            page.apply_redactions()
            # "helv"(Helvetica)는 한글 글리프가 없어 한글이 물음표로 깨진다 — 내장 CJK
            # 폰트인 "korea-s"를 쓴다. 라틴 문자는 자간이 다소 벌어지지만(CID 폰트 특성)
            # 한글이 통째로 깨지거나 사라지는 것보단 훨씬 낫다.
            _insert_textbox_fit(page, rect, new_text, font_size)
            applied += section_applied

        buf = io.BytesIO()
        doc.save(buf)
    finally:
        doc.close()
    return buf.getvalue(), applied, skipped


def _apply_text(
    path: Path, accepted: list[ProofreadFindingRead], sections: dict[int, DocumentSection]
) -> tuple[bytes, int, list[str]]:
    text = path.read_text(encoding="utf-8", errors="replace")

    bounds: list[tuple[int, int]] = []
    pos = 0
    for m in BLANK_LINE_RE.finditer(text):
        bounds.append((pos, m.start()))
        pos = m.end()
    bounds.append((pos, len(text)))
    blocks = [(s, e) for s, e in bounds if text[s:e].strip()]

    applied = 0
    skipped: list[str] = []
    by_section = _group_by_section(accepted, skipped)

    edits: list[tuple[int, int, str]] = []
    for section_id, section_findings in by_section.items():
        section = sections.get(section_id)
        if section is None:
            skipped.extend(f"finding {f.id}: section {section_id} 없음" for f in section_findings)
            continue
        if section.order_no >= len(blocks):
            skipped.extend(
                f"finding {f.id}: order {section.order_no} 범위 밖" for f in section_findings
            )
            continue
        start, end = blocks[section.order_no]
        raw_block = text[start:end]

        if section.section_type == "heading":
            current = raw_block.strip().lstrip("#").strip()
        else:
            current = raw_block.strip()
        if current != _expected_text(section).strip():
            skipped.extend(
                f"finding {f.id}: 블록 텍스트가 분석 시점과 달라짐" for f in section_findings
            )
            continue

        new_block = raw_block
        section_applied = 0
        for f in section_findings:
            if f.original in new_block:
                new_block = new_block.replace(f.original, f.suggestion, 1)
                section_applied += 1
            else:
                skipped.append(f"finding {f.id}: 원문 '{f.original}'을 찾지 못함")
        if section_applied == 0:
            continue
        edits.append((start, end, new_block))
        applied += section_applied

    edits.sort(key=lambda e: e[0])
    out_parts: list[str] = []
    cursor = 0
    for start, end, new_block in edits:
        out_parts.append(text[cursor:start])
        out_parts.append(new_block)
        cursor = end
    out_parts.append(text[cursor:])
    return "".join(out_parts).encode("utf-8"), applied, skipped


_APPLY_DISPATCH = {
    "docx": _apply_docx,
    "hwpx": _apply_hwpx,
    "pptx": _apply_pptx,
    "pdf": _apply_pdf,
    "txt": _apply_text,
    "md": _apply_text,
}


class CorrectionService:
    def __init__(
        self,
        documents: DocumentService,
        analyses: AiAnalysisResultRepository,
        versions: DocumentVersionRepository,
        storage: StorageProvider,
    ) -> None:
        self.documents = documents
        self.analyses = analyses
        self.versions = versions
        self.storage = storage

    def apply(self, document_id: int, analysis_id: int) -> ApplyOutcome:
        record = self.analyses.get(analysis_id)
        if record is None or record.document_id != document_id:
            raise AnalysisNotFoundError(analysis_id)
        if record.status != "COMPLETED":
            raise AnalysisNotCompletedError(analysis_id)

        result = record.result or {}
        if result.get("applied_version_id"):
            raise AlreadyAppliedError(analysis_id)

        findings = rehydrate_findings(result)
        accepted = [f for f in findings if f.status == FindingStatus.ACCEPTED]
        if not accepted:
            raise NoAcceptedFindingsError(analysis_id)

        src_version = self.versions.get(record.document_version_id)
        if src_version is None:
            raise VersionNotFoundError(record.document_version_id)
        ext = src_version.original_file_name.rsplit(".", 1)[-1].lower()
        apply_fn = _APPLY_DISPATCH.get(ext)
        if apply_fn is None:
            raise UnsupportedFileTypeError(ext)

        sections = {s.id: s for s in self.versions.list_sections(src_version.id)}
        src_path = self.storage.resolve(src_version.storage_path)

        new_bytes, applied, skipped = apply_fn(src_path, accepted, sections)
        if applied == 0:
            raise NoAcceptedFindingsError(analysis_id)

        new_version = self.documents.create_version(
            document_id,
            filename=src_version.original_file_name,
            content_type=src_version.mime_type,
            data=new_bytes,
        )

        result["applied_version_id"] = new_version.id
        result["applied_at"] = datetime.now(timezone.utc).isoformat()
        record.result = result
        flag_modified(record, "result")
        self.analyses.update(record)

        return ApplyOutcome(new_version, applied, len(skipped), skipped[:_MAX_SKIPPED_REASONS])


def get_correction_service(
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
    documents: DocumentService = Depends(get_document_service),
) -> CorrectionService:
    storage = LocalStorageProvider(settings.storage_root)
    return CorrectionService(
        documents=documents,
        analyses=AiAnalysisResultRepository(db),
        versions=DocumentVersionRepository(db),
        storage=storage,
    )
